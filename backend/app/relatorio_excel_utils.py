from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

_ROXO = "6A1B9A"
_ROXO_CLARO = "F3E5F5"


def _estilizar_cabecalho(ws, linha, num_colunas):
    for col in range(1, num_colunas + 1):
        celula = ws.cell(row=linha, column=col)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(start_color=_ROXO, end_color=_ROXO, fill_type="solid")
        celula.alignment = Alignment(horizontal="center")


def _autoajustar_colunas(ws):
    for coluna in ws.columns:
        maior = 0
        letra = get_column_letter(coluna[0].column)
        for celula in coluna:
            if celula.value:
                maior = max(maior, len(str(celula.value)))
        ws.column_dimensions[letra].width = maior + 4


def gerar_relatorio_excel(nome_evento: str, dados: dict) -> bytes:
    wb = Workbook()

    # --- Aba: Resumo ---
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo["A1"] = "Relatório Financeiro - Evently"
    ws_resumo["A1"].font = Font(bold=True, size=14, color=_ROXO)
    ws_resumo["A2"] = nome_evento
    ws_resumo["A2"].font = Font(bold=True, size=12)
    ws_resumo["A3"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}"

    ws_resumo["A5"] = "Total Vendido"
    ws_resumo["B5"] = dados["total_vendido"]
    ws_resumo["B5"].number_format = 'R$ #,##0.00'
    ws_resumo["A6"] = "Número de Transações"
    ws_resumo["B6"] = dados["numero_transacoes"]
    ws_resumo["A5"].font = Font(bold=True)
    ws_resumo["A6"].font = Font(bold=True)
    _autoajustar_colunas(ws_resumo)

    # --- Aba: Produtos Mais Vendidos ---
    ws_produtos = wb.create_sheet("Produtos Mais Vendidos")
    ws_produtos.append(["Produto", "Quantidade", "Valor Total"])
    _estilizar_cabecalho(ws_produtos, 1, 3)
    for p in dados["produtos_mais_vendidos"]:
        ws_produtos.append([p["nome"], p["quantidade"], p["valor_total"]])
    for linha in range(2, ws_produtos.max_row + 1):
        ws_produtos.cell(row=linha, column=3).number_format = 'R$ #,##0.00'
    _autoajustar_colunas(ws_produtos)

    # --- Aba: Vendas por Barraca (já ordenado da que mais vendeu) ---
    ws_barracas = wb.create_sheet("Vendas por Barraca")
    ws_barracas.append(["Barraca", "Valor Total"])
    _estilizar_cabecalho(ws_barracas, 1, 2)
    for b in dados["vendas_por_barraca"]:
        ws_barracas.append([b["barraca"], b["valor_total"]])
    for linha in range(2, ws_barracas.max_row + 1):
        ws_barracas.cell(row=linha, column=2).number_format = 'R$ #,##0.00'
    _autoajustar_colunas(ws_barracas)

    # --- Aba: Vendas por Período (dia + hora) ---
    ws_periodo = wb.create_sheet("Vendas por Período")
    ws_periodo.append(["Período", "Valor Total"])
    _estilizar_cabecalho(ws_periodo, 1, 2)
    for p in dados["vendas_por_periodo"]:
        ws_periodo.append([p["periodo"], p["valor_total"]])
    for linha in range(2, ws_periodo.max_row + 1):
        ws_periodo.cell(row=linha, column=2).number_format = 'R$ #,##0.00'
    _autoajustar_colunas(ws_periodo)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()